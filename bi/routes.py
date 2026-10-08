from __future__ import annotations

from flask import jsonify, request, abort, render_template, send_file, session
from io import BytesIO
from html import escape
import re
import secrets
from datetime import datetime
from sqlalchemy import update, delete

from .filters import BIValidationError, parse_filters
from .services import catalog_payload, filter_options, overview, executive
from .departments import department, SECTIONS
from .explore import explore
from .builder import (BIDashboard, WIDGETS, BACKGROUNDS, ACCENTS, GRID_COLUMNS,
                      CHART_STYLES, STYLES, ensure_table, validate_layout,
                      serialize, can_access, dataset, drill_records, focus_field)
from .semantic import (FIELDS as SEMANTIC_FIELDS, catalog_payload as semantic_catalog_payload,
                       visual as semantic_visual, drill as semantic_drill)


def register_bi_routes(
    app,
    login_required,
    require_permission,
    get_cdb,
    get_current_company,
    get_current_user,
    get_company_by_id,
    has_permission,
    today_ist,
):
    """Register the Phase-1 canonical BI API.

    Existing /api/bi/dashboard remains untouched for backward compatibility.
    Phase 2 can migrate the visual dashboard onto these versioned endpoints.
    """

    record_permission = has_permission

    def has_permission(permission, action='view'):
        # BI access does not grant access to a source workspace. Apply the
        # source entitlement to existing Explorer, Builder and summary queries.
        access = app.extensions.get('module_access')
        source_module = ('orderflow' if permission in ('stock', 'purchase_orders', 'sales_orders', 'delivery_challans')
                         else 'hr' if permission.startswith('hr') else 'core')
        return record_permission(permission, action) and (not access or access(source_module))

    def _context():
        company_id = get_current_company()
        if not company_id:
            raise BIValidationError("No active company is selected.")
        company = get_company_by_id(company_id)
        if company is None:
            raise BIValidationError("Active company could not be found.")
        cdb = get_cdb()
        if cdb is None:
            raise BIValidationError("Could not connect to the active company database.")
        base_currency = (getattr(company, "currency", None) or "INR").upper()
        return company_id, company, cdb, base_currency

    from .business_routes import register_business_routes
    register_business_routes(app, login_required, require_permission, get_cdb,
        get_current_company, get_current_user, get_company_by_id, has_permission, today_ist)

    from .warehouse_routes import register_warehouse_routes
    register_warehouse_routes(app, login_required, require_permission, get_cdb,
                              get_current_company, get_current_user, get_company_by_id)

    from .predictive_routes import register_predictive_routes
    register_predictive_routes(app, login_required, require_permission, get_cdb,
                               get_current_company, get_current_user, get_company_by_id)

    @app.route("/api/bi/v1/health", methods=["GET"], endpoint="api_bi_v1_health")
    @login_required
    @require_permission("analytics", "view")
    def api_bi_v1_health():
        try:
            company_id, company, _cdb, base_currency = _context()
            return jsonify({
                "ok": True,
                "version": "1.0",
                "company_id": company_id,
                "company_name": company.company_name,
                "base_currency": base_currency,
                "canonical_sales_source": "customer_invoices",
                "canonical_purchase_source": "purchase_invoices",
                "branch_filter": False,
            })
        except BIValidationError as exc:
            return jsonify({"ok": False, "error": str(exc)}), 400

    @app.route("/api/bi/v1/catalog", methods=["GET"], endpoint="api_bi_v1_catalog")
    @login_required
    @require_permission("analytics", "view")
    def api_bi_v1_catalog():
        return jsonify(catalog_payload())

    @app.route("/api/bi/v1/filter-options", methods=["GET"], endpoint="api_bi_v1_filter_options")
    @login_required
    @require_permission("analytics", "view")
    def api_bi_v1_filter_options():
        try:
            company_id, _company, cdb, base_currency = _context()
            return jsonify(filter_options(cdb, company_id, has_permission, base_currency))
        except BIValidationError as exc:
            return jsonify({"error": str(exc)}), 400

    @app.route("/api/bi/v1/overview", methods=["GET"], endpoint="api_bi_v1_overview")
    @login_required
    @require_permission("analytics", "view")
    def api_bi_v1_overview():
        try:
            company_id, _company, cdb, base_currency = _context()
            today = today_ist()
            filters = parse_filters(request.args, company_id, today)
            return jsonify(overview(cdb, filters, has_permission, base_currency, today))
        except BIValidationError as exc:
            return jsonify({"error": str(exc)}), 400

    @app.route("/api/bi/v1/executive", methods=["GET"], endpoint="api_bi_v1_executive")
    @login_required
    @require_permission("analytics", "view")
    def api_bi_v1_executive():
        # Phase 2 exposes company-wide values. Row-level employee/branch scope is
        # not present in the Phase 1 engine; allow owners only until it exists.
        if get_current_user().get("role") not in ("owner", "super_admin"):
            abort(403)
        try:
            company_id, _company, cdb, base_currency = _context()
            filters = parse_filters(request.args, company_id, today_ist())
            if any((filters.employee_id, filters.country, filters.client_id, filters.supplier_id,
                    filters.product_category, filters.expense_category)):
                raise BIValidationError("The executive dashboard currently supports date filtering only.")
            return jsonify(executive(cdb, filters, has_permission, base_currency, today_ist()))
        except BIValidationError as exc:
            return jsonify({"error": str(exc)}), 400

    @app.route('/bi-departments', endpoint='bi_departments')
    @login_required
    @require_permission('analytics', 'view')
    def bi_departments():
        if get_current_user().get('role') not in ('owner', 'super_admin'):
            abort(403)
        company_id, company, _cdb, _currency = _context()
        today = today_ist().isoformat()
        return render_template('bi_departments.html', company=company, active='bi_intelligence',
                               today_str=today, base_currency=_currency)

    @app.route('/api/bi/v1/departments/<section>', endpoint='api_bi_v1_department')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_v1_department(section):
        if get_current_user().get('role') not in ('owner', 'super_admin'):
            abort(403)
        if section not in SECTIONS:
            abort(404)
        try:
            company_id, _company, cdb, base_currency = _context()
            filters = parse_filters(request.args, company_id, today_ist())
            if any((filters.employee_id, filters.country, filters.client_id, filters.supplier_id,
                    filters.product_category, filters.expense_category)):
                raise BIValidationError('Phase 3 department views currently support date filtering only.')
            return jsonify(department(cdb, filters, has_permission, base_currency, today_ist(), section))
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400

    def _explore_request(limit):
        company_id, _company, cdb, base_currency = _context()
        # Country is a drill-path key here, whereas Phase 1 uses the same name
        # for a global filter. Parse only dates before applying explorer levels.
        date_args = {key: request.args[key] for key in ('from_date', 'to_date') if key in request.args}
        filters = parse_filters(date_args, company_id, today_ist())
        if any(request.args.get(key) for key in ('employee_id', 'client_id', 'supplier_id',
                                                 'product_category', 'expense_category', 'branch')):
            raise BIValidationError('Use the explorer drill path for dimensions; only dates are global filters.')
        return explore(cdb, filters, request.args, has_permission, base_currency, limit=limit)

    @app.route('/bi-explore', endpoint='bi_explore')
    @login_required
    @require_permission('analytics', 'view')
    def bi_explore():
        if get_current_user().get('role') not in ('owner', 'super_admin'):
            abort(403)
        _company_id, company, _cdb, _currency = _context()
        return render_template('bi_explore.html', company=company, active='bi_intelligence')

    @app.route('/api/bi/v1/explore', endpoint='api_bi_v1_explore')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_v1_explore():
        if get_current_user().get('role') not in ('owner', 'super_admin'):
            abort(403)
        try:
            return jsonify(_explore_request(100))
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400

    @app.route('/api/bi/v1/explore/export/<format>', endpoint='api_bi_v1_explore_export')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_v1_explore_export(format):
        if get_current_user().get('role') not in ('owner', 'super_admin'):
            abort(403)
        if format not in ('xlsx', 'pdf'):
            abort(404)
        try:
            data = _explore_request(5000)
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400
        title = 'Qiyadah BI - ' + data['level'].title()
        rows = data['rows']
        if format == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = 'BI Detail'
            ws.append([title])
            ws.append(['Level', data['level'], 'Currency', data['currency']])
            ws.append(['Label', 'Value', 'Records'])
            for row in rows:
                label = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(row['label']))
                if label.lstrip().startswith(('=', '+', '-', '@')):
                    label = "'" + label
                ws.append([label, row['value'], row['count']])
            ws.append(['Total across all groups', data['total'], data['row_count']])
            if data['truncated']:
                ws.append(['TRUNCATED: only the first 5000 rows are exported.'])
            for note in data['warnings']:
                ws.append([note])
            for cell in ws[3]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='0C4840')
            ws.column_dimensions['A'].width = 58
            ws.column_dimensions['B'].width = 19
            ws.column_dimensions['C'].width = 18
            for row in ws.iter_rows(min_row=4, max_col=2):
                if row[1].data_type == 'n':
                    row[1].number_format = '#,##0.00'
            output = BytesIO()
            wb.save(output)
            output.seek(0)
            return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             as_attachment=True, download_name='qiyadah_bi_detail.xlsx')
        from xhtml2pdf import pisa
        body = ''.join('<tr><td>%s</td><td align="right">%s</td><td align="right">%s</td></tr>' %
                       (escape(str(r['label'])), f"{r['value']:,.2f}", r['count']) for r in rows)
        notes = ''.join('<p>%s</p>' % escape(n) for n in data['warnings'])
        if data['truncated']:
            notes += '<p>Truncated: first 5000 rows shown.</p>'
        html = ('<html><head><style>@page{size:A4;margin:18mm}body{font-family:Helvetica;font-size:9pt}'
                'td,th{border-bottom:1px solid #ddd;padding:6px}table{width:100%}th{background-color:#e5eee9}'
                '</style></head><body><h1>%s</h1><p>Currency: %s</p><table><thead><tr>'
                '<th align="left">Label</th><th align="right">Value</th><th align="right">Records</th>'
                '</tr></thead><tbody>%s</tbody></table><p>Total across all groups: %.2f</p>%s</body></html>') % (
                    escape(title), escape(data['currency']), body, data['total'], notes)
        output = BytesIO()
        status = pisa.CreatePDF(html, dest=output, encoding='UTF-8')
        if status.err:
            return jsonify({'error': 'PDF rendering failed.'}), 500
        output.seek(0)
        return send_file(output, mimetype='application/pdf', as_attachment=True,
                         download_name='qiyadah_bi_detail.pdf')

    def _builder_access():
        user = get_current_user()
        if user.get('role') not in ('owner', 'super_admin') and not record_permission('analytics', 'view'):
            abort(403)
        company_id, company, cdb, currency = _context()
        return user, company_id, company, cdb, currency

    def _builder_csrf():
        expected = session.get('bi_builder_csrf')
        supplied = request.headers.get('X-BI-CSRF')
        if not expected or not supplied or not secrets.compare_digest(expected, supplied):
            abort(403)

    @app.route('/bi-builder', endpoint='bi_builder')
    @login_required
    @require_permission('analytics', 'view')
    def bi_builder():
        user, company_id, company, cdb, currency = _builder_access()
        ensure_table(cdb)
        token = session.get('bi_builder_csrf') or secrets.token_urlsafe(32)
        session['bi_builder_csrf'] = token
        return render_template('bi_builder.html', company=company, active='bi_intelligence',
                               bi_csrf=token, bi_semantic_catalog=semantic_catalog_payload())

    @app.route('/api/bi/v1/builder/catalog', endpoint='api_bi_builder_catalog')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_catalog():
        _builder_access()
        return jsonify({'widgets': {kind: list(value['keys']) for kind, value in WIDGETS.items()},
                        'chart_styles': {key: list(value) for key, value in CHART_STYLES.items()},
                        'widget_styles': {key: list(value) for key, value in STYLES.items()},
                        'backgrounds': list(BACKGROUNDS), 'accents': list(ACCENTS),
                        'visibility': ['private', 'owner_team'], 'grid_columns': GRID_COLUMNS,
                        'max_widgets': 30, 'semantic': semantic_catalog_payload()})

    def _semantic_request(drilling=False):
        user, company_id, company, cdb, currency = _builder_access()
        allowed = {'from_date', 'to_date', 'employee_id', 'country', 'client_id',
                   'supplier_id', 'product_category', 'expense_category', *SEMANTIC_FIELDS}
        if drilling:
            allowed.add('group')
        if set(request.args) - allowed:
            raise BIValidationError('Unsupported chart field or filter.')
        filters = parse_filters(request.args, company_id, today_ist())
        focus_field(filters)
        spec = {key: request.args[key] for key in SEMANTIC_FIELDS if key in request.args}
        if 'semantic_measures' in spec:
            spec['semantic_measures'] = request.args.getlist('semantic_measures')
        if 'semantic_limit' in spec:
            try:
                spec['semantic_limit'] = int(spec['semantic_limit'])
            except ValueError:
                raise BIValidationError('Choose a valid top-N limit.') from None
        if drilling:
            return semantic_drill(cdb, filters, has_permission, currency, spec, request.args.get('group'))
        return semantic_visual(cdb, filters, has_permission, currency, spec)

    @app.route('/api/bi/v1/builder/visual', endpoint='api_bi_builder_visual')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_visual():
        try:
            return jsonify(_semantic_request())
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400

    @app.route('/api/bi/v1/builder/visual/drill', endpoint='api_bi_builder_visual_drill')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_visual_drill():
        try:
            return jsonify(_semantic_request(drilling=True))
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400

    @app.route('/api/bi/v1/builder/data', endpoint='api_bi_builder_data')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_data():
        user, company_id, company, cdb, currency = _builder_access()
        try:
            allowed = {'from_date', 'to_date', 'employee_id', 'country', 'client_id',
                       'supplier_id', 'product_category', 'expense_category'}
            if set(request.args) - allowed:
                raise BIValidationError('Unsupported dashboard filter.')
            filters = parse_filters(request.args, company_id, today_ist())
            focus_field(filters)
            return jsonify(dataset(cdb, filters, has_permission, currency, today_ist()))
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400

    @app.route('/api/bi/v1/builder/drill', endpoint='api_bi_builder_drill')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_drill():
        user, company_id, company, cdb, currency = _builder_access()
        try:
            allowed = {'from_date', 'to_date', 'employee_id', 'country', 'client_id',
                       'supplier_id', 'product_category', 'expense_category',
                       'chart', 'group', 'series', 'period'}
            if set(request.args) - allowed:
                raise BIValidationError('Unsupported drill filter.')
            filters = parse_filters(request.args, company_id, today_ist())
            return jsonify(drill_records(cdb, filters, has_permission, currency,
                                         request.args.get('chart'), request.args.get('group'),
                                         request.args.get('series'), request.args.get('period')))
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400

    @app.route('/api/bi/v1/builder/dashboards', methods=['GET', 'POST'], endpoint='api_bi_builder_dashboards')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_dashboards():
        user, company_id, company, cdb, currency = _builder_access()
        ensure_table(cdb)
        if request.method == 'GET':
            rows = cdb.query(BIDashboard).filter_by(company_id=company_id).order_by(BIDashboard.updated_at.desc()).all()
            return jsonify({'dashboards': [serialize(row) for row in rows if can_access(row, company_id, user)]})
        _builder_csrf()
        if request.content_length is not None and request.content_length > 40000:
            return jsonify({'error': 'Dashboard request is too large.'}), 413
        try:
            title, visibility, layout = validate_layout(request.get_json(silent=True))
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400
        row = BIDashboard(company_id=company_id, owner_user_id=user.get('user_id'),
                          title=title, visibility=visibility, layout_json=layout,
                          revision=1, created_at=datetime.utcnow(), updated_at=datetime.utcnow())
        cdb.add(row)
        cdb.commit()
        return jsonify(serialize(row)), 201

    @app.route('/api/bi/v1/builder/dashboards/<int:dashboard_id>', methods=['GET', 'PUT', 'DELETE'], endpoint='api_bi_builder_dashboard')
    @login_required
    @require_permission('analytics', 'view')
    def api_bi_builder_dashboard(dashboard_id):
        user, company_id, company, cdb, currency = _builder_access()
        ensure_table(cdb)
        row = cdb.query(BIDashboard).filter_by(id=dashboard_id, company_id=company_id).first()
        if not row or not can_access(row, company_id, user):
            abort(404)
        if request.method == 'GET':
            return jsonify(serialize(row))
        if row.owner_user_id != user.get('user_id') and user.get('role') != 'super_admin':
            abort(403)
        _builder_csrf()
        data = request.get_json(silent=True)
        if not isinstance(data, dict) or type(data.get('revision')) is not int:
            return jsonify({'error': 'Current dashboard revision is required.'}), 400
        if data['revision'] != row.revision:
            return jsonify({'error': 'This dashboard changed in another session. Reload before saving.', 'revision': row.revision}), 409
        if request.method == 'DELETE':
            deleted = cdb.execute(delete(BIDashboard).where(BIDashboard.id == dashboard_id,
                BIDashboard.company_id == company_id, BIDashboard.revision == row.revision)).rowcount
            if not deleted:
                cdb.rollback()
                return jsonify({'error': 'Dashboard changed in another session. Reload before deleting.'}), 409
            cdb.commit()
            return jsonify({'deleted': True})
        if request.content_length is not None and request.content_length > 40000:
            return jsonify({'error': 'Dashboard request is too large.'}), 413
        try:
            title, visibility, layout = validate_layout(data)
        except BIValidationError as exc:
            return jsonify({'error': str(exc)}), 400
        changed = cdb.execute(update(BIDashboard).where(BIDashboard.id == dashboard_id,
            BIDashboard.company_id == company_id, BIDashboard.revision == data['revision']).values(
            title=title, visibility=visibility, layout_json=layout,
            revision=data['revision'] + 1, updated_at=datetime.utcnow())).rowcount
        if not changed:
            cdb.rollback()
            return jsonify({'error': 'Dashboard changed in another session. Reload before saving.'}), 409
        cdb.commit()
        cdb.expire_all()
        row = cdb.query(BIDashboard).filter_by(id=dashboard_id, company_id=company_id).first()
        return jsonify(serialize(row))

        # SQL/database errors are intentionally not swallowed here. Flask's
        # existing OperationalError handler should continue to surface/recover
        # real connection failures consistently with the rest of Qiyadah.
